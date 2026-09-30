"""Iteration 2 - Location fields, master options, empty DB, import/export.

Covers: MapView removed, Alamat & Lokasi fields added (village, district,
city_regency, province, postal_code, country, latitude, longitude,
location_accuracy, location_updated_at, location_updated_by).

NOTE: pytest is configured to run with xdist (loadscope). We keep all mutating
tests in a single class so they run on one worker and can share CREATED_IDS.
"""
import io
import csv
import os
import time
import pytest
import requests
from openpyxl import load_workbook

# Resolve base URL from frontend .env (public URL)
_env_val = os.environ.get("EXPO_PUBLIC_BACKEND_URL")
if not _env_val:
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("EXPO_PUBLIC_BACKEND_URL="):
                _env_val = ln.split("=", 1)[1].strip()
                break
BASE = (_env_val or "").rstrip("/")


def _session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


def _admin_token(s):
    r = s.post(f"{BASE}/api/admin/login",
               json={"username": "admin", "password": "admin123"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def _user_token(s):
    r = s.post(f"{BASE}/api/login",
               json={"username": "user", "password": "user123"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def AH(t):
    return {"Authorization": f"Bearer {t}"}


EXPECTED_HEADERS = [
    "Customer Code", "Customer Name", "Segment", "Purchasing Size", "Area",
    "Status", "Bad Debt", "Bad Debt Nominal", "Address",
    "Village", "District", "City/Regency", "Province", "Postal Code", "Country",
    "Latitude", "Longitude", "Location Accuracy",
    "Phone", "WhatsApp", "PIC Name", "Payment Terms", "Credit Limit",
]


# ---- Auth (standalone) ------------------------------------------------------
class TestAuth:
    def test_user_login(self):
        s = _session()
        r = s.post(f"{BASE}/api/login",
                   json={"username": "user", "password": "user123"})
        assert r.status_code == 200 and "access_token" in r.json()

    def test_admin_login(self):
        s = _session()
        r = s.post(f"{BASE}/api/admin/login",
                   json={"username": "admin", "password": "admin123"})
        assert r.status_code == 200 and "access_token" in r.json()

    def test_bad_creds(self):
        s = _session()
        r = s.post(f"{BASE}/api/login",
                   json={"username": "user", "password": "wrong"})
        assert r.status_code in (400, 401)


# ---- Master options endpoint shape -----------------------------------------
class TestMasterOptions:
    def test_master_options_keys(self):
        s = _session()
        t = _admin_token(s)
        r = s.get(f"{BASE}/api/admin/master-options", headers=AH(t))
        assert r.status_code == 200
        j = r.json()
        for k in ("segment", "purchasing_size", "area", "top"):
            assert isinstance(j.get(k), list)


# ---- Main flow (single class => single xdist worker) -----------------------
class TestLocationFlow:
    CREATED: list = []
    MASTER_IDS: dict = {}

    @classmethod
    def setup_class(cls):
        cls.s = _session()
        cls.admin = _admin_token(cls.s)
        cls.user = _user_token(cls.s)
        # Seed a few master values so customer creation passes min_length=1
        for et, name in (("segment", "TEST_SegA"),
                         ("purchasing_size", "TEST_PSA"),
                         ("area", "TEST_AreaA"),
                         ("top", "TEST_TopA")):
            r = cls.s.post(f"{BASE}/api/admin/master/{et}",
                           json={"name": name, "description": ""},
                           headers=AH(cls.admin))
            if r.status_code in (200, 201) and isinstance(r.json(), dict):
                cls.MASTER_IDS[et] = (r.json().get("id"), name)
            else:
                cls.MASTER_IDS[et] = (None, name)

    @classmethod
    def teardown_class(cls):
        # remove created customers
        for cid in cls.CREATED:
            cls.s.delete(f"{BASE}/api/customers/{cid}", headers=AH(cls.admin))
        # remove created master rows
        for et, (mid, _n) in cls.MASTER_IDS.items():
            if mid:
                cls.s.delete(f"{BASE}/api/admin/master/{et}/{mid}",
                             headers=AH(cls.admin))

    def _payload(self, code):
        return {
            "customer_code": code,
            "customer_name": f"TEST_LOC_{code}",
            "segment": self.MASTER_IDS["segment"][1],
            "purchasing_size": self.MASTER_IDS["purchasing_size"][1],
            "area": self.MASTER_IDS["area"][1],
            "status": "Active",
            "phone": "081200000000",
            "whatsapp": "081200000000",
            "pic_name": "TEST PIC",
            "payment_terms": self.MASTER_IDS["top"][1],
            "credit_limit": 0,
            "bad_debt": False,
            "bad_debt_nominal": 0,
            "address": "Jl. Uji No. 1",
            "village": "TEST_VillageZ",
            "district": "TEST_DistrictZ",
            "city_regency": "TEST_CityZ",
            "province": "TEST_ProvinceZ",
            "postal_code": "12345",
            "country": "Indonesia",
            "latitude": -6.2000,
            "longitude": 106.8166,
            "location_accuracy": 12.5,
        }

    def test_01_create_persists_all_location_fields(self):
        p = self._payload("TEST_LOC_A1")
        r = self.s.post(f"{BASE}/api/customers", json=p, headers=AH(self.admin))
        assert r.status_code == 201, r.text
        j = r.json()
        self.__class__.CREATED.append(j["id"])
        for f in ("village", "district", "city_regency", "province",
                  "postal_code", "country", "location_accuracy",
                  "latitude", "longitude"):
            assert j.get(f) == p[f], f"{f}: {j.get(f)!r} != {p[f]!r}"
        assert j.get("location_updated_at"), "auto-stamp missing"
        assert j.get("location_updated_by") == "admin"

    def test_02_get_returns_all_fields(self):
        cid = self.CREATED[0]
        r = self.s.get(f"{BASE}/api/customers/{cid}", headers=AH(self.user))
        assert r.status_code == 200
        j = r.json()
        assert j["village"] == "TEST_VillageZ"
        assert j["province"] == "TEST_ProvinceZ"
        assert j["location_updated_by"] == "admin"

    def test_03_update_no_coord_change_keeps_audit(self):
        cid = self.CREATED[0]
        prev = self.s.get(f"{BASE}/api/customers/{cid}", headers=AH(self.admin)).json()
        p = self._payload("TEST_LOC_A1")
        p["pic_name"] = "TEST PIC EDITED"
        time.sleep(1.1)
        r = self.s.put(f"{BASE}/api/customers/{cid}", json=p, headers=AH(self.admin))
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["location_updated_at"] == prev["location_updated_at"]
        assert j["location_updated_by"] == prev["location_updated_by"]
        assert j["pic_name"] == "TEST PIC EDITED"

    def test_04_update_coord_change_restamps(self):
        cid = self.CREATED[0]
        prev = self.s.get(f"{BASE}/api/customers/{cid}", headers=AH(self.admin)).json()
        p = self._payload("TEST_LOC_A1")
        p["latitude"] = -6.3
        p["longitude"] = 106.9
        time.sleep(1.1)
        r = self.s.put(f"{BASE}/api/customers/{cid}", json=p, headers=AH(self.admin))
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["latitude"] == -6.3
        assert j["location_updated_at"] != prev["location_updated_at"]

    def test_05_search_village(self):
        r = self.s.get(f"{BASE}/api/customers?search=TEST_VillageZ",
                       headers=AH(self.user))
        assert r.status_code == 200
        assert any(c.get("id") in self.CREATED for c in r.json())

    def test_06_search_city(self):
        r = self.s.get(f"{BASE}/api/customers?search=TEST_CityZ",
                       headers=AH(self.user))
        assert r.status_code == 200
        assert any(c.get("id") in self.CREATED for c in r.json())

    def test_07_search_province(self):
        r = self.s.get(f"{BASE}/api/customers?search=TEST_ProvinceZ",
                       headers=AH(self.user))
        assert r.status_code == 200
        assert any(c.get("id") in self.CREATED for c in r.json())

    def test_08_search_postal(self):
        r = self.s.get(f"{BASE}/api/customers?search=12345",
                       headers=AH(self.user))
        assert r.status_code == 200
        assert any(c.get("id") in self.CREATED for c in r.json())

    def test_09_csv_export_header(self):
        r = self.s.get(f"{BASE}/api/admin/export", headers=AH(self.admin))
        assert r.status_code == 200
        j = r.json()
        assert "csv" in j, f"expected csv key, got {list(j.keys())}"
        first = j["csv"].splitlines()[0]
        got = next(csv.reader(io.StringIO(first)))
        assert got == EXPECTED_HEADERS, f"Header mismatch:\nexp={EXPECTED_HEADERS}\ngot={got}"
        # Also verify our created row round-trips new columns
        rows = list(csv.reader(io.StringIO(j["csv"])))
        our = [r for r in rows[1:] if r and r[0] == "TEST_LOC_A1"]
        assert our, "created row not in export"
        row = our[0]
        # Position of Village etc.
        idx = EXPECTED_HEADERS.index
        assert row[idx("Village")] == "TEST_VillageZ"
        assert row[idx("City/Regency")] == "TEST_CityZ"
        assert row[idx("Province")] == "TEST_ProvinceZ"
        assert row[idx("Postal Code")] == "12345"
        assert row[idx("Country")] == "Indonesia"

    def test_10_xlsx_template_headers(self):
        r = self.s.get(f"{BASE}/api/admin/template.xlsx", headers=AH(self.admin))
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        for col in ("Village", "District", "City/Regency", "Province",
                    "Postal Code", "Country", "Location Accuracy"):
            assert col in headers, f"missing {col}"

    def test_11_xlsx_export_headers(self):
        r = self.s.get(f"{BASE}/api/admin/export.xlsx", headers=AH(self.admin))
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert headers[:len(EXPECTED_HEADERS)] == EXPECTED_HEADERS

    def test_12_import_preview_accepts_new_columns(self):
        row = {
            "Customer Code": "TEST_LOC_IMP_1",
            "Customer Name": "TEST Import Loc",
            "Segment": self.MASTER_IDS["segment"][1],
            "Purchasing Size": self.MASTER_IDS["purchasing_size"][1],
            "Area": self.MASTER_IDS["area"][1],
            "Status": "Active", "Bad Debt": "No", "Bad Debt Nominal": 0,
            "Address": "Jl Imp", "Village": "V", "District": "D",
            "City/Regency": "C", "Province": "P", "Postal Code": "99999",
            "Country": "Indonesia",
            "Latitude": -6.1, "Longitude": 106.7, "Location Accuracy": 5,
            "Phone": "0812", "WhatsApp": "0812", "PIC Name": "PIC",
            "Payment Terms": self.MASTER_IDS["top"][1], "Credit Limit": 0,
        }
        r = self.s.post(f"{BASE}/api/admin/import/preview",
                        json={"rows": [row]}, headers=AH(self.admin))
        assert r.status_code == 200, r.text
        j = r.json()
        counts = j.get("counts", {})
        assert counts.get("total") == 1, j
        assert counts.get("error", 0) == 0, f"import preview errors: {j}"
